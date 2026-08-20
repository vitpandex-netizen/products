"""
Генератор/обновление файла тикеров UZSE в формате, найденном в
~/Downloads/UZSE_tickers_final_27_02_2026.xlsx

Колонки (подтверждено чтением реального файла):
  ticker    - тикер на UZSE (напр. URTS, SQBN, HMKB)
  amount    - размер лота (в найденном файле везде 100)
  price     - последняя известная цена
  SearchName- строка для поиска новостей: "(RU-имя OR Lat-имя OR EN-имя OR TICKER)"
  Date      - дата снимка цены, формат DD.MM.YYYY
  Time      - время снимка, формат HH:MM

Использование:
  from shared.tools.uzse_ticker_file import read_tickers, write_tickers, update_price

  rows = read_tickers("UZSE_tickers_final_27_02_2026.xlsx")
  update_price(rows, "SQBN", price=33.5, date="28.02.2026", time="18:00")
  write_tickers(rows, "UZSE_tickers_new.xlsx")
"""

from dataclasses import dataclass, asdict
from typing import List, Optional
import openpyxl

COLUMNS = ["ticker", "amount", "price", "SearchName", "Date", "Time"]


@dataclass
class TickerRow:
    ticker: str
    amount: Optional[float]
    price: Optional[float]
    SearchName: str
    Date: str
    Time: str


def read_tickers(path: str) -> List[TickerRow]:
    """Прочитать xlsx в этом формате. Пропускает заголовок."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header = [str(h).strip() if h else "" for h in rows[0]]
    result = []
    for r in rows[1:]:
        if not r or not r[0]:
            continue
        data = dict(zip(header, r))
        result.append(TickerRow(
            ticker=data.get("ticker", ""),
            amount=data.get("amount"),
            price=data.get("price") or data.get("buy price"),
            SearchName=data.get("SearchName", ""),
            Date=data.get("Date", ""),
            Time=data.get("Time", ""),
        ))
    return result


def write_tickers(rows: List[TickerRow], path: str) -> None:
    """Записать список TickerRow в xlsx с теми же колонками, что в оригинале."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "UZSE Tickers"
    ws.append(COLUMNS)
    for row in rows:
        ws.append([row.ticker, row.amount, row.price, row.SearchName, row.Date, row.Time])
    wb.save(path)


def update_price(rows: List[TickerRow], ticker: str, price: float, date: str, time: str) -> bool:
    """Обновить цену конкретного тикера в списке (in-place). Вернёт False если тикер не найден."""
    for row in rows:
        if row.ticker == ticker:
            row.price = price
            row.Date = date
            row.Time = time
            return True
    return False


def build_search_query(ticker: str, names: List[str]) -> str:
    """Собрать SearchName в формате оригинала: "(Имя1 OR Имя2 OR TICKER)" """
    parts = names + [ticker]
    return "(" + " OR ".join(parts) + ")"


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("Usage: python uzse_ticker_file.py <path_to_xlsx>")
        sys.exit(1)
    rows = read_tickers(sys.argv[1])
    print(f"Загружено {len(rows)} тикеров:")
    for r in rows[:5]:
        print(f"  {r.ticker}: {r.price} ({r.Date} {r.Time})")
